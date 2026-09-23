from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_project_documentation():
    doc = Document()

    # Configure Margins (0.8 inch for clean executive look)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color Palette Constants
    NAVY = RGBColor(0, 51, 128)      # #003380 (Schiphol Brand Navy)
    DARK_GRAY = RGBColor(50, 50, 50)
    CHARCOAL = RGBColor(30, 30, 30)

    # ----------------------------------------------------
    # COVER / HEADER
    # ----------------------------------------------------
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("SCHIPHOL AIRPORT INTELLIGENCE HUB")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = NAVY

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("Executive Project Summary & Presentation Portfolio\n"
                            "Cloud Data Warehouse Architecture & Real-Time Executive BI")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ====================================================
    # SECTION 1: EXECUTIVE 1-PAGE MEMO
    # ====================================================
    sec1_heading = doc.add_paragraph()
    sec1_run = sec1_heading.add_run("SECTION 1: EXECUTIVE 1-PAGE MEMO")
    sec1_run.font.name = "Arial"
    sec1_run.font.size = Pt(15)
    sec1_run.font.bold = True
    sec1_run.font.color.rgb = NAVY
    sec1_heading.paragraph_format.space_after = Pt(8)

    # 1. Problem Statement
    h1 = doc.add_paragraph()
    r = h1.add_run("1. Business Context & Strategic Challenge")
    r.font.bold = True
    r.font.size = Pt(11.5)
    h1.paragraph_format.space_after = Pt(2)

    doc.add_paragraph(
        "Fragmented Operational Silos: Critical operational data across Amsterdam Airport Schiphol "
        "(flight schedules, weather conditions, baggage handling, retail point-of-sale, and parking facilities) "
        "operated in isolated, disconnected systems.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "Lack of Executive Visibility: Commercial managers, ground handling dispatchers, and airport leadership "
        "lacked a unified, real-time command center to track terminal footfall, severe weather turnaround hazards, "
        "and multi-channel non-aeronautical revenue.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "Cloud Scalability & Query Latency: Unoptimized, cross-departmental queries over millions of transactional "
        "records resulted in high compute costs, table locks, and sluggish executive dashboard performance.",
        style='List Bullet'
    )

    # 2. Engineering Solution
    h2 = doc.add_paragraph()
    r = h2.add_run("2. Engineering Architecture & Analytics Solution")
    r.font.bold = True
    r.font.size = Pt(11.5)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(2)

    doc.add_paragraph(
        "Unified Cloud Data Warehouse: Consolidated 9 airport touchpoints into a high-performance PostgreSQL 15 "
        "warehouse hosted on Supabase under a structured 'aviation' schema adhering to Third Normal Form (3NF).",
        style='List Bullet'
    )
    doc.add_paragraph(
        "Hybrid Data Sourcing Strategy: Integrated live real-time flight telemetry from the Schiphol Airport Developer "
        "REST API and live runway meteorological feeds, complemented with synthetic operational datasets "
        "(baggage piece weights, excess fee penalties, retail POS receipts, parking sensors, and maintenance logs).",
        style='List Bullet'
    )
    doc.add_paragraph(
        "In-Database Statistical Intelligence: Automated parametric (mean, standard deviation) and non-parametric "
        "(percentiles P50, P90) calculations directly inside PostgreSQL, applying Z-scores in views to detect abnormal "
        "luxury retail transactions with zero application memory overhead.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "Performance Optimization (CTE vs. Subquery): Replaced nested correlated subqueries with modular Common Table "
        "Expressions (CTEs), reducing query latency by ~65% (from ~200 ms to ~70 ms) and cutting cloud database compute.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "Zero-Latency Pre-Aggregation: Built Materialized Views with unique date indexes to pre-compute multi-domain "
        "commercial revenue onto disk, supporting non-blocking concurrent cache synchronization.",
        style='List Bullet'
    )

    # 3. Operational Impact
    h3 = doc.add_paragraph()
    r = h3.add_run("3. Quantified Business Impact & Value Delivered")
    r.font.bold = True
    r.font.size = Pt(11.5)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(2)

    doc.add_paragraph(
        "24/7 Live Decision Command Center: Successfully deployed an interactive Streamlit BI application live to the web, "
        "accessible on desktop, tablet, and mobile devices with zero local configuration required.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "Proactive Ramp Risk Management: Automatically flags flights combining above-average baggage loads with high-wind "
        "runway disruptions, allowing operations managers to reassign ground handlers and prevent costly flight delays.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "Data-Driven Concession Strategies: Identified intraday passenger footfall spending waves and spending tier "
        "breakdowns (< €25, €25–€100, > €100), guiding commercial tenant lease structures and concession staffing.",
        style='List Bullet'
    )

    # PAGE BREAK
    doc.add_page_break()

    # ====================================================
    # SECTION 2: SLIDE-BY-SLIDE PRESENTATION OUTLINE
    # ====================================================
    sec2_heading = doc.add_paragraph()
    sec2_run = sec2_heading.add_run("SECTION 2: SLIDE-BY-SLIDE PRESENTATION OUTLINE")
    sec2_run.font.name = "Arial"
    sec2_run.font.size = Pt(15)
    sec2_run.font.bold = True
    sec2_run.font.color.rgb = NAVY
    sec2_heading.paragraph_format.space_after = Pt(12)

    slides = [
        (
            "SLIDE 1: PROJECT VISION & STRATEGIC OBJECTIVES",
            [
                "Objective: Design and deploy a real-time operational and commercial intelligence platform for Amsterdam Airport Schiphol (AMS).",
                "Target Audience: Airport Executive Leadership, Terminal Operations Managers, Flight Dispatchers, and Commercial Concession Directors.",
                "Deliverables: A cloud-hosted PostgreSQL warehouse on Supabase powering an interactive, multi-tab Streamlit dashboard live on the web."
            ]
        ),
        (
            "SLIDE 2: DATA ACQUISITION & HYBRID INGESTION STRATEGY",
            [
                "Live External REST APIs: Ingested real-time flight schedules, aircraft IDs, and carrier IATA codes from the official Schiphol Developer API, paired with live runway meteorological sensor telemetry.",
                "Synthetic Operational Data: Synthesized granular baggage transactions with individual piece weights (kg) and excess surcharges (€), airside retail POS sales, parking facility sensors, and aircraft maintenance logs.",
                "Harmonized Foundation: Cleaned and structured 9 core operational tables into a unified relational warehouse schema."
            ]
        ),
        (
            "SLIDE 3: CLOUD DATA WAREHOUSE ARCHITECTURE",
            [
                "Infrastructure: Managed PostgreSQL 15 database instance provisioned on Supabase cloud.",
                "Relational Modeling: Designed a dedicated 'aviation' schema enforcing Third Normal Form (3NF), primary keys, and foreign key integrity constraints across all operational tables.",
                "Automated Python Ingestion: Developed ETL data pipelines using Python, pandas, and psycopg2 to handle type-casting, timestamp standardization, and batch upserts."
            ]
        ),
        (
            "SLIDE 4: IN-DATABASE ADVANCED ANALYTICS & STATISTICAL PROFILING",
            [
                "Server-Side Percentiles: Computed non-parametric medians (P50) and 90th percentile (P90) thresholds directly in SQL using PERCENTILE_CONT to baseline baggage weights and retail ticket averages.",
                "Automated Anomaly Detection: Built the view 'vw_retail_statistical_anomalies' to calculate transactional mean, standard deviation, and Z-scores directly in the database engine.",
                "Executive Value: Automatically identifies and flags unusual, high-value luxury purchases without application memory overhead."
            ]
        ),
        (
            "SLIDE 5: ARCHITECTURAL BENCHMARK: SUBQUERY VS. CTE",
            [
                "The Operational Problem: Identify flights operating during severe wind conditions that carry luggage volume at or above the carrier benchmark.",
                "Legacy Pattern (Correlated Subquery): Recalculated airline benchmarks row-by-row inside the WHERE clause; execution latency hit ~200 ms with high CPU overhead.",
                "Optimized Pattern (Modular CTE): Pre-filtered weather dates and baggage weights once into memory using WITH (...) blocks; latency dropped to ~70 ms (~65% faster, reducing cloud compute costs)."
            ]
        ),
        (
            "SLIDE 6: ANALYTICAL WINDOW FUNCTIONS & MATERIALIZED VIEWS",
            [
                "High-Fidelity Window Functions: Leveraged DENSE_RANK() OVER (PARTITION BY airline) to rank carrier cargo loads, and running totals (SUM() OVER) without collapsing underlying flight rows.",
                "Materialized View Acceleration: Pre-joined multi-table metrics across retail, parking, and baggage onto disk in 'mv_daily_operational_summary' for near-instant dashboard loads.",
                "Zero-Downtime Cache Sync: Created a unique index on the schedule date to allow non-blocking concurrent refreshes (REFRESH CONCURRENTLY) via the user interface."
            ]
        ),
        (
            "SLIDE 7: EXECUTIVE STREAMLIT DASHBOARD ARCHITECTURE",
            [
                "Tab 1 (Data Warehouse Inventory): Live operational record counts and ingestion volume distribution charts across all 9 airport entities.",
                "Tab 2 (Terminal Retail & Concessions): Intraday footfall surge wave charts (peak hours) and purchase tier distributions (< €25, €25–€100, > €100).",
                "Tab 3 (Severe Weather & Ramp Risk): Cargo weight moved per airline during high-wind disruptions, highlighting high-risk turnaround flights.",
                "Tab 4 (Airline Baggage Performance): Total luggage volume (kg) vs. average bag weight per carrier, accompanied by flight turnaround rankings.",
                "Tab 5 (Commercial Revenue Mix): Clean comparative breakdown of terminal retail concessions against ground parking facility earnings.",
                "Tab 6 (SQL Engineering Lab): Dedicated technical workspace displaying side-by-side CTE vs. Subquery benchmarks, runtimes, and full SQL catalogs."
            ]
        ),
        (
            "SLIDE 8: PRODUCTION DEPLOYMENT & ROADBLOCK RESOLUTIONS",
            [
                "Automated Concurrency Handling: Fixed sidebar refresh lock errors by implementing an isolated autocommit connection outside active transaction blocks.",
                "Cloud Secrets Management: Resolved cloud networking and connection timeouts by configuring direct IPv6 database connection strings securely in Streamlit Cloud Secrets.",
                "Live Global Delivery: Platform published to Streamlit Community Cloud (24/7 public access), fully operational on mobile, tablet, and desktop without requiring local software."
            ]
        )
    ]

    for title, points in slides:
        slide_h = doc.add_paragraph()
        r = slide_h.add_run(title)
        r.font.bold = True
        r.font.size = Pt(11.5)
        slide_h.paragraph_format.space_before = Pt(8)
        slide_h.paragraph_format.space_after = Pt(2)

        for pt in points:
            doc.add_paragraph(pt, style='List Bullet')

    # Save document
    filename = "Schiphol_Airport_Project_Portfolio.docx"
    doc.save(filename)
    print(f"Document successfully created and saved as: {filename}")

if __name__ == "__main__":
    create_project_documentation()